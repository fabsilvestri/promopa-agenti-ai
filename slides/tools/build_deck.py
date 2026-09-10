"""Costruisce la lezione a partire da Template_RSTLess.pptx.

Regola dura: massimo 3 bullet per slide. Se servono di più, la slide si
sdoppia con titolo "Titolo 1/n".
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
from PIL import Image
import copy

TPL = "slides/tools/Template_RSTLess.pptx"
OUT = "slides/Lezione_15_settembre_Agenti_AI_PromoPA.pptx"
FIG = "slides/tools/fig/"
FOTO = "slides/tools/fig/foto/"
# Il template RSTLess e' bordeaux. I titoli passano a un blu notte: il filetto
# del footer e il logo restano bordeaux e fanno da accento, non da protagonista.
BORD = RGBColor(0x82, 0x24, 0x33)
BLU = RGBColor(0x1B, 0x3A, 0x5F)
NERO = RGBColor(0x21, 0x21, 0x21); GRIG = RGBColor(0x59, 0x59, 0x59)
AZZ = RGBColor(0xE9, 0xF1, 0xF8); VERDE = RGBColor(0x2C, 0x6E, 0x49)
BIANCO = RGBColor(0xFF, 0xFF, 0xFF); CHIARO = RGBColor(0xEC, 0xEC, 0xEC)
VELO = RGBColor(0x1A, 0x0A, 0x0D)  # quasi nero con una punta di bordeaux
FOOTER = "Promo PA Fondazione  |  Agenti AI per i processi interni  |  15.09.2026"

prs = Presentation(TPL)
W, H = prs.slide_width, prs.slide_height  # 9144000 x 5143500

# --- footer nel master
for shp in prs.slide_master.shapes:
    if shp.has_text_frame and "Project name" in shp.text_frame.text:
        r = shp.text_frame.paragraphs[0].runs[0]
        r.text = FOOTER

# --- elimina le slide di esempio
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

L_TITLE, L_SECTION, L_BODY, L_TWO, L_TITLEONLY, L_BLANK = [prs.slide_layouts[i] for i in range(6)]

BULLET_COUNT_MAX = 3


def _style_title(ph, size=27):
    tf = ph.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = BLU; r.font.name = "Arial"


def _bullets(tf, items, size=19, color=NERO, space=10):
    """Scrive bullet veri (buChar), uno per paragrafo. Sottoelenchi con tuple (testo, [sub])."""
    assert len(items) <= BULLET_COUNT_MAX, f"troppi bullet: {items}"
    tf.word_wrap = True
    first = True
    for it in items:
        sub = None
        if isinstance(it, tuple):
            it, sub = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _set_bullet(p, it, size, color, space, level=0)
        if sub:
            for s in sub:
                p2 = tf.add_paragraph()
                _set_bullet(p2, s, size - 3, GRIG, space - 4, level=1)


def _set_bullet(p, text, size, color, space, level=0):
    p.text = ""
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Arial"
    p.space_after = Pt(space)
    pPr = p._p.get_or_add_pPr()
    marL = 285750 if level == 0 else 571500
    pPr.set("marL", str(marL)); pPr.set("indent", str(-228600)); pPr.set("lvl", str(level))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr")); c = etree.SubElement(buClr, qn("a:srgbClr")); c.set("val", "1B3A5F")
    buFont = etree.SubElement(pPr, qn("a:buFont")); buFont.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•" if level == 0 else "–")


def _textbox(slide, x, y, w, h, text, size=14, color=NERO, bold=False, align=PP_ALIGN.LEFT, font="Arial", fill=None, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(90000); tf.margin_top = tf.margin_bottom = Emu(60000)
    if fill is not None:
        tb.fill.solid(); tb.fill.fore_color.rgb = fill
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font; r.font.italic = italic
    return tb


def _code(slide, x, y, w, h, code, size=11):
    tb = _textbox(slide, x, y, w, h, code, size=size, color=NERO, font="Courier New", fill=AZZ)
    tb.line.color.rgb = RGBColor(0xC9, 0xD6, 0xE3)
    return tb


def _dietro(slide, shape, posizione):
    """Sposta una forma in fondo alla pila, sotto i segnaposto del layout."""
    el = shape._element
    el.getparent().remove(el)
    slide.shapes._spTree.insert(posizione, el)


def _foto_riempi(slide, percorso, x, y, w, h):
    """Foto ritagliata per riempire il rettangolo, senza deformarla.

    PowerPoint stira l'immagine dentro la cornice: se prima ritagliamo alla
    stessa proporzione della cornice, lo stiramento e' nullo e le persone non
    sembrano alte il doppio.
    """
    pic = slide.shapes.add_picture(percorso, Emu(x), Emu(y), width=Emu(w), height=Emu(h))
    with Image.open(percorso) as im:
        nativo = im.size[0] / im.size[1]
    cornice = w / h
    if nativo > cornice:
        quota = (1 - cornice / nativo) / 2
        pic.crop_left = pic.crop_right = quota
    else:
        quota = (1 - nativo / cornice) / 2
        pic.crop_top = pic.crop_bottom = quota
    return pic


def _velo(slide, x, y, w, h, colore=VELO, opacita=52):
    """Rettangolo semitrasparente sopra la foto, perche' il testo si legga."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    box.line.fill.background()
    box.shadow.inherit = False
    box.fill.solid()
    box.fill.fore_color.rgb = colore
    srgb = box._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", str(int(opacita * 1000)))
    return box


def _sfondo_fotografico(slide, foto, x, y, w, h, opacita=52):
    """Foto a tutta larghezza con velo, mandata dietro ai segnaposto."""
    pic = _foto_riempi(slide, foto, x, y, w, h)
    velo = _velo(slide, x, y, w, h, opacita=opacita)
    _dietro(slide, pic, 2)
    _dietro(slide, velo, 3)
    return pic


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def title_slide(title, subtitle, notes="", foto=None):
    s = prs.slides.add_slide(L_TITLE)
    # La fascia azzurra del template va da 564050 a 3630050: la foto la copre.
    if foto:
        _sfondo_fotografico(s, foto, 0, 564050, W, 3066000, opacita=55)
    colore_titolo = BIANCO if foto else BLU
    colore_sotto = CHIARO if foto else GRIG
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = colore_titolo; r.font.name = "Arial"
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = subtitle
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(16); r.font.color.rgb = colore_sotto; r.font.name = "Arial"
    # rimuove i placeholder inutilizzati (foto e testi dei relatori)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx in (2, 3, 4, 5):
            ph._element.getparent().remove(ph._element)
    _textbox(s, 2700000, 3835000, 6200000, 420000, "Fabrizio Silvestri  |  Sapienza Università di Roma, DIAG", size=14, color=GRIG, anchor=MSO_ANCHOR.MIDDLE)
    _notes(s, notes)
    return s


def section(title, subtitle="", notes="", foto=None):
    s = prs.slides.add_slide(L_SECTION)
    # Foto a tutta pagina fino al filetto del footer, che resta leggibile.
    if foto:
        _sfondo_fotografico(s, foto, 0, 0, W, 4500000, opacita=52)
    colore_titolo = BIANCO if foto else BLU
    colore_sotto = CHIARO if foto else GRIG
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = colore_titolo; r.font.name = "Arial"
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = subtitle
            for p in ph.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                pPr = p._p.get_or_add_pPr()
                for e in list(pPr):
                    if e.tag in (qn("a:buChar"), qn("a:buFont"), qn("a:buClr"), qn("a:buNone")):
                        pPr.remove(e)
                etree.SubElement(pPr, qn("a:buNone"))
                pPr.set("marL", "0"); pPr.set("indent", "0")
                for r in p.runs:
                    r.font.size = Pt(16); r.font.color.rgb = colore_sotto; r.font.name = "Arial"
    _notes(s, notes)
    return s


def bullets_foto(title, items, foto, notes="", size=19, foto_frac=0.40):
    """Bullet a sinistra, fotografia a destra. Per le slide con spazio vuoto."""
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    left_w = int(8784000 * (1 - foto_frac)) - 90000
    tb = s.shapes.add_textbox(Emu(180000), Emu(1200000), Emu(left_w), Emu(3100000))
    tb.text_frame.margin_left = Emu(0)
    _bullets(tb.text_frame, items, size=size, space=12)
    fx = 180000 + left_w + 180000
    fw = 8964000 - fx
    _foto_riempi(s, foto, fx, 1250000, fw, int(fw / 1.5))
    _notes(s, notes)
    return s


def _etichetta(slide, cx, cy, testo, size=13, fondo=BIANCO, colore=BLU, opacita=92):
    """Targhetta di testo vero sopra la fotografia.

    Le immagini generate non sanno scrivere: la disposizione la da' la foto,
    le parole le mette PowerPoint. Cosi' restano leggibili e correggibili.
    """
    righe = testo.split("\n")
    larghezza = int(max(len(r) for r in righe) * size * 6900) + 200000
    altezza = int(len(righe) * size * 15200) + 110000
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Emu(int(cx - larghezza / 2)), Emu(int(cy - altezza / 2)),
                                 Emu(larghezza), Emu(altezza))
    box.adjustments[0] = 0.18
    box.fill.solid(); box.fill.fore_color.rgb = fondo
    srgb = box._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    alpha = etree.SubElement(srgb, qn("a:alpha")); alpha.set("val", str(int(opacita * 1000)))
    box.line.color.rgb = colore
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(60000)
    tf.margin_top = tf.margin_bottom = Emu(20000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, riga in enumerate(righe):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = PP_ALIGN.CENTER
        r = par.add_run(); r.text = riga
        r.font.size = Pt(size if i == 0 else size - 2)
        r.font.bold = (i == 0)
        r.font.color.rgb = colore
        r.font.name = "Arial"
    return box


def _freccia(slide, x1, y1, x2, y2, colore=BLU, spessore=2.0):
    """Connettore diritto con la punta, disegnato sopra la fotografia."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
    c.line.color.rgb = colore
    c.line.width = Pt(spessore)
    linea = c._element.spPr.find(qn("a:ln"))
    testa = etree.SubElement(linea, qn("a:tailEnd"))
    testa.set("type", "triangle"); testa.set("w", "med"); testa.set("len", "med")
    return c


def schema_foto(title, foto, etichette, notes="", caption="", frecce=(),
                alto=1160000, basso=3860000, size=13, velo=0):
    """Schema costruito su una fotografia: la foto da' la disposizione,
    le etichette e le frecce sono oggetti veri sopra di essa.

    etichette: sequenza di (testo, x, y) con x e y relativi al riquadro
    della foto, da 0 a 1. frecce: sequenza di (x1, y1, x2, y2), stesse unita'.
    """
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    x0, w = 180000, 8784000
    h = basso - alto
    _foto_riempi(s, foto, x0, alto, w, h)
    if velo:
        _velo(s, x0, alto, w, h, colore=BIANCO, opacita=velo)
    for x1, y1, x2, y2 in frecce:
        _freccia(s, x0 + x1 * w, alto + y1 * h, x0 + x2 * w, alto + y2 * h)
    for testo, rx, ry in etichette:
        _etichetta(s, x0 + rx * w, alto + ry * h, testo, size=size)
    if caption:
        _textbox(s, 180000, basso + 80000, 8784000, 300000, caption, size=12, color=GRIG, italic=True)
    _notes(s, notes)
    return s


def foto_piena(title, foto, caption="", notes=""):
    """Titolo in alto, fotografia sotto a tutta larghezza. Serve per parlarci sopra."""
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    alto, basso = 1140000, 4160000
    _foto_riempi(s, foto, 180000, alto, 8784000, basso - alto)
    if caption:
        _textbox(s, 180000, basso + 60000, 8784000, 300000, caption, size=11, color=GRIG, italic=True)
    _notes(s, notes)
    return s


def bullets(title, items, notes="", size=22):
    s = prs.slides.add_slide(L_BODY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    body = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
    body.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    body.top = Emu(1250000)
    _bullets(body.text_frame, items, size=size, space=14)
    _notes(s, notes)
    return s


def bullets_image(title, items, image, notes="", img_w_frac=0.52, size=19):
    """Testo a sinistra, immagine a destra."""
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    left_w = int(8784000 * (1 - img_w_frac)) - 90000
    tb = s.shapes.add_textbox(Emu(180000), Emu(1134300), Emu(left_w), Emu(3400000))
    tb.text_frame.margin_left = Emu(0)
    _bullets(tb.text_frame, items, size=size)
    ix = 180000 + left_w + 180000
    iw = 8964000 - ix
    pic = s.shapes.add_picture(image, Emu(ix), Emu(1180000), width=Emu(iw))
    if pic.height > 3400000:
        ratio = 3400000 / pic.height
        pic.height = Emu(3400000); pic.width = Emu(int(pic.width * ratio))
        pic.left = Emu(ix + (iw - pic.width) // 2)
    _notes(s, notes)
    return s


def image_full(title, image, caption="", notes="", max_h=3000000):
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    pic = s.shapes.add_picture(image, Emu(180000), Emu(1150000), width=Emu(8784000))
    if pic.height > max_h:
        ratio = max_h / pic.height
        pic.height = Emu(max_h); pic.width = Emu(int(pic.width * ratio))
        pic.left = Emu((W - pic.width) // 2)
    if caption:
        _textbox(s, 180000, 4230000, 8784000, 300000, caption, size=11, color=GRIG, italic=True)
    _notes(s, notes)
    return s


BASSO_UTILE = 4300000  # sotto questa quota comincia l'aria prima del footer


def image_top_bullets(title, image, items, notes="", img_h=1500000, size=19):
    """Figura larga in alto, bullet sotto.

    L'altezza chiesta viene ridotta se i bullet non ci starebbero: meglio una
    figura un po' piu' piccola che una riga di testo sopra il filetto.
    """
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    righe_stimate = sum(1 + len(str(i)) // 88 for i in items)
    alto_testo = int(righe_stimate * (size * 1.35 + 8) * 12700)
    img_h = min(img_h, BASSO_UTILE - alto_testo - 150000 - 1120000)
    pic = s.shapes.add_picture(image, Emu(180000), Emu(1120000), height=Emu(img_h))
    if pic.width > 8784000:
        ratio = 8784000 / pic.width
        pic.width = Emu(8784000); pic.height = Emu(int(pic.height * ratio))
    pic.left = Emu((W - pic.width) // 2)
    y = 1120000 + pic.height + 150000
    tb = s.shapes.add_textbox(Emu(180000), Emu(y), Emu(8784000), Emu(4600000 - y))
    tb.text_frame.margin_left = Emu(0)
    _bullets(tb.text_frame, items, size=size, space=8)
    _notes(s, notes)
    return s


def bullets_code(title, items, code, notes="", code_frac=0.55, size=17, code_size=10.5):
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    left_w = int(8784000 * (1 - code_frac)) - 90000
    tb = s.shapes.add_textbox(Emu(180000), Emu(1134300), Emu(left_w), Emu(3400000))
    tb.text_frame.margin_left = Emu(0)
    _bullets(tb.text_frame, items, size=size)
    cx = 180000 + left_w + 180000
    _code(s, cx, 1150000, 8964000 - cx, 3200000, code, size=code_size)
    _notes(s, notes)
    return s


def code_full(title, code, notes="", size=11):
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    _code(s, 180000, 1150000, 8784000, 3200000, code, size=size)
    _notes(s, notes)
    return s


def numbers(title, stats, items, notes=""):
    """Tre numeri grandi in alto, tre bullet sotto."""
    s = prs.slides.add_slide(L_TITLEONLY)
    s.shapes.title.text = title; _style_title(s.shapes.title)
    n = len(stats); gap = 180000; cw = (8784000 - gap * (n - 1)) // n
    for i, (num, lab) in enumerate(stats):
        x = 180000 + i * (cw + gap)
        box = s.shapes.add_shape(1, Emu(x), Emu(1150000), Emu(cw), Emu(1350000))
        box.fill.solid(); box.fill.fore_color.rgb = AZZ; box.line.fill.background()
        box.shadow.inherit = False
        _textbox(s, x, 1180000, cw, 750000, num, size=44, color=BLU, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _textbox(s, x, 1930000, cw, 520000, lab, size=13, color=GRIG, align=PP_ALIGN.CENTER)
    tb = s.shapes.add_textbox(Emu(180000), Emu(2700000), Emu(8784000), Emu(1900000))
    tb.text_frame.margin_left = Emu(0)
    _bullets(tb.text_frame, items, size=18)
    _notes(s, notes)
    return s


# =====================================================================
# CONTENUTO
# =====================================================================

title_slide("Agenti AI per i processi interni",
            "Dal caso helpdesk formazione a un pilota misurabile\nPromo PA Fondazione, 15 settembre 2026",
            foto=FOTO + "copertina.jpg",
            notes="Le fotografie di questo deck sono generate con AI: dirlo qui, in dieci secondi. Non e' un adempimento, e' buona educazione, ed e' l'esempio piu' corto che avete di cosa vuol dire dichiararlo. Seconda lezione dopo quella del 22 giugno su ChatGPT.")

# ---------- 1. Riaggancio ----------
section("1. Riaggancio e tesi della giornata", "10 minuti", foto=FOTO + "sez1_riaggancio.jpg",
        notes="Chiedere: chi ha usato Progetti o GPT personalizzati da giugno? Due minuti di giro di tavolo, non di piu'.")

bullets_foto("Da giugno a oggi", [
    "A giugno: ChatGPT come collega. Prompt, Progetti, documenti, RAG",
    "Oggi: dal collega al processo. Cosa succede quando l'AI entra in un flusso di lavoro",
    "Caso di studio: helpdesk formazione",
], FOTO + "c03_da_giugno.jpg",
    notes="Richiamare che ChatGPT Business è già in uso in Fondazione: tutto ciò che vedremo si può provare lì, senza acquistare niente.")

bullets_foto("Che cosa sapete già fare", [
    "Scrivere un prompt che dice il compito, il formato e a chi è destinato",
    "Tenere i documenti in un Progetto, così il contesto non si riscrive ogni volta",
    "Accorgervi quando il modello inventa, e chiedergli su cosa si basa",
], FOTO + "ap_portate.jpg",
    notes="Serve a dire: non ripartiamo da zero. Queste tre cose bastano per la demo di oggi. Se qualcuno non le ha provate, si metta in coppia con chi le ha provate.")

image_top_bullets("La scaletta di oggi", FIG + "scaletta.png", [
    "Prima il vocabolario, poi il vostro caso, poi i limiti",
    "L'ultima mezz'ora è vostra: si lavora in gruppi su un processo della Fondazione",
    "Molte slide sono immagini: interrompete quando volete, le domande valgono di più",
], img_h=1950000,
    notes="Dichiarare il patto: si può interrompere. Le slide sono tante ma molte durano dieci secondi.")

image_top_bullets("Chatbot, workflow, agente", FIG + "chatbot_workflow_agente.png", [
    "Un chatbot risponde. Non agisce",
    "Un workflow segue passi decisi da noi; l'AI sta dentro alcuni passi",
    "Un agente decide da solo i passi. Serve raramente, costa controllo",
], img_h=1950000,
    notes="Tesi provocatoria ma vera: la maggior parte dei 'processi ripetitivi' vuole un workflow con AI dentro, non un agente autonomo. L'autonomia si compra con la prevedibilità.")

foto_piena("L'agente autonomo che vi immaginate", FOTO + "gag_robot.jpg",
           caption="Autonomo davvero: decide da solo, e non arriva alla tastiera.",
           notes="Slide per ridere, e per fissare un concetto: quando un fornitore dice 'agente autonomo', chiedete cosa sa fare da solo davvero. Spesso la risposta è: poco, e sotto sorveglianza.")

bullets_foto("I tre tipi, con un esempio ciascuno", [
    "Chatbot: chiedete a ChatGPT come si scrive una convocazione. Risponde, e finisce lì",
    "Workflow: la richiesta arriva, il modello la classifica, una regola la smista, un umano invia",
    "Agente: gli dite 'sistema l'helpdesk' e decide lui i passi. Oggi non lo vogliamo",
], FOTO + "b08_tre_tipi.jpg", size=18,
    notes="Il secondo è quello che costruiremo. Vale la pena dirlo tre volte nella giornata.")

bullets_foto("La tesi in tre righe", [
    "Prima si scompone il processo, poi si decide cosa automatizzare",
    "Dove serve giudizio o responsabilità, resta un punto di controllo umano",
    "Si misura prima di automatizzare, e si rimisura dopo",
], FOTO + "c05_tesi.jpg",
    notes="Queste tre righe tornano alla fine come 'tre errori da evitare', rovesciate.")

# ---------- 2. Anatomia ----------
section("2. Anatomia di un sistema agentico", "20 minuti", foto=FOTO + "sez2_anatomia.jpg",
        notes="Vocabolario minimo per parlare con fornitori e consulenti senza farsi vendere un agente quando serve una regola.")

bullets_image("I cinque pezzi 1/2", [
    "Modello: legge testo, produce testo o dati strutturati. Non sa nulla di voi",
    "Strumenti: e-mail, Zoom, Moodle, sito. Il modello li usa, non li possiede",
    "Memoria: tassonomia, storico, regole scritte da voi",
], FIG + "anatomia_agente.png", img_w_frac=0.50,
    notes="Il modello è la parte più famosa e la meno importante dal punto di vista del processo: cambia ogni sei mesi. Tassonomia e strumenti restano.")

bullets_image("I cinque pezzi 2/2", [
    "Ciclo di azione: leggi, decidi, agisci, verifica. Poi ricomincia",
    "Punto di controllo umano: soglia di confidenza, coda da verificare, invio manuale",
    "Un sistema senza punto di controllo non è coraggioso, è incompleto",
], FIG + "anatomia_agente.png", img_w_frac=0.50,
    notes="Il punto di controllo è progettato, non aggiunto dopo. Se il fornitore non sa dirvi dove sta, la risposta è 'da nessuna parte'.")

bullets_foto("Il modello: cosa sa e cosa non sa", [
    "Sa la lingua, il mondo generico, la forma dei documenti amministrativi",
    "Non sa i vostri corsi, i vostri operatori, le vostre scadenze, i vostri clienti",
    "Tutto quello che non sa glielo dovete passare a ogni chiamata: non se lo ricorda",
], FOTO + "an_modello.jpg",
    notes="Esempio concreto: il modello non sa che 'Bilancio enti locali' è un vostro corso. Lo impara dal file della tassonomia, ogni volta, non una volta per sempre.")

bullets_foto("Gli strumenti: il modello non possiede niente", [
    "Legge la casella se qualcuno gliela collega. Non ha una casella",
    "Ogni strumento è un permesso che concedete, e che potete togliere in un minuto",
    "Nel nostro disegno il modello legge e scrive bozze. Non invia, non cancella",
], FOTO + "an_strumenti.jpg",
    notes="Domanda utile in aula: quali permessi dareste, e quali no? Di solito la risposta è: leggere sì, scrivere forse, inviare mai.")

bullets_foto("La memoria: la tassonomia siete voi", [
    "La memoria che conta non è il ricordo della conversazione: è un file che scrivete voi",
    "Tassonomia, regole, esempi difficili: stanno in un file, non dentro il prompt",
    "Cambiare un corso è cambiare una riga, non riscrivere il sistema",
], FOTO + "an_memoria.jpg",
    notes="Questo è il punto che rende il sistema vostro e non del fornitore. Chi possiede il file possiede il comportamento.")

image_top_bullets("Il ciclo di azione", FIG + "ciclo_azione.png", [
    "Quattro passi, e poi da capo. È tutto qui",
    "Il passo che sparisce sempre nelle offerte è il quarto: verifica",
    "Senza verifica non è un ciclo, è una freccia",
], img_h=1950000,
    notes="Nel nostro caso la verifica è doppia: la soglia di confidenza automatica, e l'operatore che rilegge la bozza prima di inviare.")

bullets_foto("Il punto di controllo umano", [
    "Si progetta prima, non si aggiunge dopo che qualcosa è andato storto",
    "Tre forme: soglia di confidenza, coda da verificare, invio manuale",
    "Nel blocco 4 vedremo che la soglia, da sola, non basta",
], FOTO + "an_controllo.jpg",
    notes="Anticipare qui il colpo di scena del blocco 4: nella nostra esecuzione reale la soglia non è mai scattata, e sedici richieste su quaranta avevano un campo sbagliato.")

bullets_image("Pattern utili 1/2", [
    "Routing: una richiesta entra, un destinatario esce. Sotto soglia, coda umana",
    "Estrazione strutturata: da testo libero a campi fissi (chi, cosa, quando)",
    "Sono i due pattern che coprono il vostro caso quasi per intero",
], FIG + "pattern.png", img_w_frac=0.58,
    notes="Anthropic e OpenAI usano nomi diversi per gli stessi pattern; l'idea è la stessa: dare al modello un compito piccolo e un formato di uscita rigido.")

bullets_foto("Pattern utili 2/2", [
    "Orchestratore e worker: un modello spezza il compito, altri lo eseguono. Utile per rendicontazioni lunghe",
    "Valutatore: un secondo passaggio controlla il primo. Utile prima dell'invio",
    "Regola d'oro: compito piccolo, formato di uscita rigido, un umano dove costa sbagliare",
], FOTO + "b19_orchestra.jpg", size=17,
    notes="Non entrare nei dettagli di orchestrazione: basta che riconoscano i nomi quando li leggono in un'offerta.")

bullets_foto("Parole da riconoscere in un'offerta 1/3", [
    "RAG: il modello cerca nei vostri documenti prima di rispondere. Utile, non magico",
    "Fine tuning: si riaddestra il modello sui vostri dati. Caro, lento, quasi mai necessario",
    "Prompt di sistema: le istruzioni fisse. È lì che vive il comportamento, chiedete di vederlo",
], FOTO + "b20_dizionario.jpg", size=16,
    notes="Se un'offerta propone fine tuning per smistare 300 richieste al mese, è un campanello. Con quei volumi non c'è nulla da riaddestrare.")

bullets_foto("Parole da riconoscere in un'offerta 2/3", [
    "Output strutturato: il modello restituisce campi fissi, non prosa. Chiedetelo sempre",
    "Guardrail: controlli che bloccano le uscite fuori regola. Vanno mostrati, non promessi",
    "Human in the loop: la persona dentro il ciclo. Chiedete in quale passo, esattamente",
], FOTO + "b21_occhiali.jpg", size=16,
    notes="'Human in the loop' senza il nome del passo è una formula vuota. Fatevi indicare la schermata che la persona vede.")

bullets_foto("Parole da riconoscere in un'offerta 3/3", [
    "Orchestrazione: un modello che ne coordina altri. Serve per compiti lunghi, non per lo smistamento",
    "Allucinazione: il modello inventa con sicurezza. Non si elimina, si contiene",
    "Token: l'unità con cui si paga. Chiedete il costo per mille richieste, non per token",
], FOTO + "b22_bilancia.jpg", size=16,
    notes="Sull'ultimo punto: nessuno sa stimare i token a occhio. Il costo per mille richieste è una domanda a cui un fornitore serio risponde subito.")

bullets_foto("Quando NON usare un agente", [
    "Se la regola si scrive in un 'se... allora', è una regola. Non serve un modello",
    "Se l'errore costa più del tempo risparmiato, serve un umano, non un agente",
    "Se non avete dati per misurarlo, non sapete se funziona. Prima i dati",
], FOTO + "c11_regola.jpg",
    notes="Esempio: 'assegna all'amministrazione tutto ciò che è fatturazione' è una regola. 'Capisci se questa e-mail parla di fatturazione' è AI.")

bullets_foto("Tre domande da fare a un fornitore", [
    "Su quali dati avete misurato l'accuratezza, e posso vedere il campione?",
    "In quale passo esatto interviene una persona, e che cosa vede sullo schermo?",
    "Se cambio la tassonomia, cosa devo toccare, quanto costa e chi lo fa?",
], FOTO + "an_fornitore.jpg",
    notes="Sono tre domande che si fanno in due minuti e che separano un fornitore serio da un venditore. Suggerire di scriverle sul quaderno.")

# ---------- 3. Il caso ----------
section("3. Il caso helpdesk formazione", "30 minuti, con demo", foto=FOTO + "sez3_helpdesk.jpg",
        notes="Dati forniti dalla Fondazione l'8 settembre: circa 300 richieste al mese, 5 caselle, nessuno strumento di tracciamento. I dodici esempi reali sono arrivati con il documento anonimizzato.")

numbers("I vostri numeri", [("~300", "richieste al mese"), ("5", "caselle e-mail diverse"), ("0", "strumenti di tracciamento")], [
    "Canali: e-mail dirette, pagine di iscrizione sul sito, Moodle",
    "Persi quasi del tutto: risposte aperte dei questionari Moodle e chat Zoom",
    "Chi risponde: chi ha la responsabilità sul tema. Funziona finché tutti ricordano tutto",
], notes="300 al mese sono 15 al giorno lavorativo. Non è volume da sistema complesso: è volume da 'una tabella e una regola di smistamento'.")

foto_piena("Quindici richieste al giorno", FOTO + "gag_carrello.jpg",
           caption="Non è un volume da piattaforma. È un volume da tabella e da regole.",
           notes="Serve a sgonfiare l'ansia: nessuno qui ha un problema di scala. Il problema è di ordine, non di volume.")

image_top_bullets("Da dove arrivano le richieste", FIG + "canali.png", [
    "Tre canali arrivano già a qualcuno: caselle, moduli, messaggi Moodle",
    "Due si perdono: questionari e chat Zoom. Sono quelli con dentro i suggerimenti",
    "Il primo lavoro non è capire il testo: è farlo arrivare tutto nello stesso posto",
], img_h=1950000,
    notes="Chiedere in aula: qualcuno rilegge le chat Zoom dopo il corso? Di solito la risposta è no, e questo apre il discorso sull'estrazione.")

foto_piena("Cinque caselle", FOTO + "gag_telefoni.jpg",
           caption="Cinque caselle e-mail sono cinque telefoni sulla stessa scrivania.",
           notes="Slide per ridere e per fissare il punto: il problema non è che le caselle siano cinque, è che nessuna sa cosa succede nelle altre quattro.")

foto_piena("Il sistema di tracciamento attuale", FOTO + "gag_postit.jpg",
           caption="Funziona benissimo, finché la persona che ha scritto i foglietti è in ufficio.",
           notes="Da dire con leggerezza, senza colpevolizzare nessuno: e' cosi' in quasi tutte le organizzazioni sotto una certa dimensione.")

bullets_foto("Cosa si perde per strada", [
    "Le chat Zoom: nessuno le rilegge, e dentro ci sono domande vere e proposte di corsi",
    "Le risposte aperte dei questionari: si contano le stelline, non si leggono le frasi",
    "Le telefonate: fuori portata per oggi, ma vale la pena contarle per un mese",
], FOTO + "hd_canali.jpg",
    notes="Il terzo punto è un compito a costo zero: un foglio accanto al telefono per quattro settimane. Spesso cambia le priorità del pilota.")

bullets_image("Le quattro esigenze", [
    "Solo due delle quattro hanno bisogno di un modello",
    "Le altre due sono integrazione e workflow: si risolvono senza AI, e vanno risolte prima",
    "Ordine consigliato: prima 1, 3 e 4, poi la 2",
], FIG + "mappa_esigenze.png", img_w_frac=0.62,
    notes="L'ordine è controintuitivo: la 2 (Zoom e Moodle) è la più affascinante e la meno urgente. Senza il contenitore unico, le richieste estratte non hanno dove andare.")

bullets_foto("Esigenza 1: repository unico", [
    "Non è AI. È un contenitore più dei connettori dalle 5 caselle e dai moduli del sito",
    "Va bene anche un foglio condiviso o un ticketing gratuito, purché sia uno solo",
    "Da qui in poi ogni richiesta ha un id, uno stato, un responsabile",
], FOTO + "c15_repository.jpg",
    notes="Menzionare strumenti a basso codice: Make, Zapier, n8n, o le automazioni di Microsoft 365 e Google Workspace che già pagate.")

foto_piena("Il contenitore che oggi non c'è", FOTO + "gag_faldoni.jpg",
           caption="Finché non c'è un posto solo, ogni richiesta trovata è una richiesta che non sa dove andare.",
           notes="E' la ragione per cui l'esigenza 1 viene prima della 2. Se estraete richieste dalle chat Zoom e non avete dove metterle, avete solo creato lavoro.")

bullets_foto("Come si costruisce il contenitore", [
    "Un foglio condiviso basta, purché sia uno solo e ogni riga abbia un id",
    "I connettori dalle caselle si fanno a basso codice, in un pomeriggio",
    "La regola vera è organizzativa: nessuno risponde più direttamente dalla propria casella",
], FOTO + "b35_imbuto.jpg", size=18,
    notes="Il terzo punto è il più difficile e non costa niente in tecnologia. È lì che i progetti falliscono.")

bullets_foto("L'integrazione è idraulica, non intelligenza", [
    "Spostare dati da un posto a un altro senza interpretarli non richiede un modello",
    "È noiosa, è la maggior parte del lavoro, e nessuno la mette nelle presentazioni",
    "Se un'offerta è tutta AI e niente integrazione, manca il 70 per cento del progetto",
], FOTO + "gag_cavi.jpg",
    notes="Nella pipeline che vedrete, tre passi su sette usano un modello. Gli altri quattro sono tubi.")

bullets_code("Esigenza 2: Zoom e Moodle", [
    "Testo non strutturato: 17 righe di chat, 4 richieste vere",
    "Il modello separa segnale e rumore e riporta le parole dell'utente",
    "Non riassume, non corregge, non inventa: solo estrae",
], """[10:03] Comune di Pisa: buongiorno
[10:11] Comune di Pisa: le slide verranno
        inviate dopo?
[10:25] Unione Valdera: non sento bene...
[10:26] Unione Valdera: ok risolto
[10:48] Comune di Livorno: domanda per il
        docente: la soglia per l'affidamento
        diretto vale anche per i servizi
        di ingegneria?
[11:41] Comune di Lucca: l'attestato viene
        rilasciato automaticamente?
[12:20] Provincia: sarebbe utile un corso
        sulla parte esecuzione del contratto
[12:29] Comune di Pisa: grazie, ottimo corso""", code_frac=0.52,
    notes="Righe da evidenziare a voce: 10:11, 10:48, 11:41, 12:20 sono richieste. 10:25 no: l'utente dice 'risolto'. 12:29 è un ringraziamento, non un feedback azionabile.")

bullets_code("Esigenza 2: cosa esce dall'estrazione", [
    "Ogni richiesta porta con sé la riga da cui viene, e il motivo",
    "Il motivo serve a voi, non al modello: è come si controlla a campione",
    "L'id lo ricostruiamo noi dall'orario, non lo lasciamo decidere al modello",
], """{
  "id_origine": "11:41",
  "mittente": "Comune di Lucca",
  "testo": "l'attestato viene rilasciato
    automaticamente o dobbiamo
    richiederlo?",
  "motivazione": "chiede una procedura,
    nessuno ha risposto in chat"
}""", code_frac=0.48,
    notes="Il terzo bullet nasce da un errore vero: alla prima esecuzione il modello metteva nell'id tutta la riga, mittente compreso, e la valutazione non trovava piu' le richieste. Ora l'orario lo estraiamo con un'espressione regolare.")

bullets_code("Esigenza 3: classificazione 1/2", [
    "La tassonomia la scrivete voi, in un file. Il modello la applica",
    "Gli operatori sono ruoli, non persone: così non invecchia",
    "Cambiare un corso non richiede di toccare il prompt",
], """# data/tassonomia.yaml
tipologie:
  - id: attestato
    descrizione: "Rilascio, correzione
      o reinvio dell'attestato"
  - id: fatturazione
    descrizione: "Fatture, CIG, split
      payment, pagamenti"
operatori:
  - id: segreteria-didattica
    competenze: [iscrizione, attestato,
                 informazioni]
  - id: amministrazione
    competenze: [fatturazione]
soglia_confidenza: 0.7""", code_frac=0.50,
    notes="Il file è nel repository. In demo cambieremo una voce e rifaremo la classificazione senza toccare il prompt.")

bullets_code("Esigenza 3: classificazione 2/2", [
    "Output strutturato: il modello restituisce questo oggetto o un errore. Mai testo libero",
    "La confidenza è dichiarata e decide se assegnare o fermarsi",
    "Sotto 0,7 la richiesta va in coda umana. È il punto di controllo",
], """{
  "corso": "bilancio-enti-locali",
  "tipologia": "reclamo",
  "urgenza": "alta",
  "operatore": "direzione",
  "confidenza": 0.89,
  "campo_meno_sicuro": "urgenza",
  "riassunto": "Pontedera: terzo sollecito
    per l'attestato, minaccia escalation
    entro venerdì"
}""", code_frac=0.48,
    notes="Esempio R007. Il reclamo prevale sull'oggetto (attestato): regola scritta nel prompt di sistema, non dedotta dal modello.")

bullets_foto("La tassonomia si scrive così", [
    "Ogni voce ha una descrizione breve: serve al modello quanto agli operatori",
    "Le voci sono poche e non si sovrappongono. Se due voci litigano, il modello sbaglia",
    "C'è sempre una voce 'altro', ed è un termometro, non una discarica",
], FOTO + "b41_caratteri.jpg", size=17,
    notes="Se 'altro' supera il dieci per cento, manca una categoria. Nel nostro caso reale manca una voce commerciale, e si vede.")

bullets_foto("Gli operatori sono ruoli, non persone", [
    "Nel file c'è 'segreteria didattica', non il nome di chi ci lavora oggi",
    "Se qualcuno cambia mansione o va in maternità, la tassonomia non invecchia",
    "Il ruolo porta con sé le competenze, e la regola di assegnazione segue da sola",
], FOTO + "hd_ruoli.jpg",
    notes="E' anche una scelta di privacy: nel file non finiscono nomi di dipendenti.")

bullets_foto("Esigenza 4: tracciamento e follow-up", [
    "Stati e scadenze sono deterministici: nuova, presa in carico, in attesa, chiusa. Nessuna AI",
    "L'AI scrive la bozza di risposta; l'operatore la rivede e la invia",
    "Segnaposto obbligatori dove il sistema non sa: [data], [importo], [link]",
], FOTO + "b43_cartellini.jpg", size=17,
    notes="La bozza con segnaposto è una scelta di progetto: preferiamo un buco visibile a una data inventata.")

image_top_bullets("Gli stati di una richiesta", FIG + "stati.png", [
    "Quattro stati e due date: presa in carico e scadenza. Niente di più",
    "Il passaggio indietro esiste: da 'in attesa utente' si torna a 'presa in carico'",
    "Qui non serve nessun modello, e infatti non ce n'è",
], img_h=1950000,
    notes="Questa è la parte che si fa in mezza giornata con un foglio condiviso, ed è quella che dà il sollievo maggiore alla segreteria.")

image_full("L'architettura, tutta insieme", FIG + "pipeline.png",
           caption="Arancio: passi con un modello. Verde: integrazione o umano. Un workflow, non un agente autonomo.",
           notes="Contare i passi con AI: tre su sette. Il resto è idraulica. È il messaggio centrale della lezione.")

bullets_foto("Demo dal vivo: cosa vedrete", [
    "Un GPT personalizzato in ChatGPT Business con la tassonomia come file di conoscenza",
    "La stessa logica in 200 righe di Python con output strutturato e valutazione",
    "La chat Zoom grezza che diventa 4 righe in tabella",
], FOTO + "c21_demo.jpg",
    notes="Sequenza in demo/README.md. Piano B senza rete: demo/output_esempio, che ora contiene l'esecuzione vera del 10 settembre. Tempo: 10 min GPT, 10 min codice, 10 min discussione.")

image_full("Demo: il risultato", FIG + "tabella_demo.png",
           caption="Dieci righe su quaranta. Righe gialle: urgenza alta. Le colonne sono quelle della tassonomia.",
           notes="Fermarsi su RE01, una richiesta vera: il partecipante ha finito il corso ma la piattaforma non gli fa scaricare l'attestato. E' un problema di attestato o di piattaforma? Il modello dice attestato, noi avevamo etichettato accesso-piattaforma. Non e' ovvio chi abbia ragione: e' esattamente il caso da portare a un umano.")

bullets_foto("Demo: le tre righe da guardare", [
    "R007: reclamo con sollecito. Il tono prevale sull'oggetto, e va alla direzione",
    "RE01: attestato bloccato dalla piattaforma. Due categorie plausibili, nessuna ovvia",
    "RE05: una notifica automatica di ordine MEPA. Non è nemmeno una domanda",
], FOTO + "b48_segnalibri.jpg", size=18,
    notes="Se resta tempo, aprire il file e cercare insieme una riga a caso. La forza della demo è che i dati sono i loro.")

code_full("Demo: una bozza di risposta (R007)", """Gentile Ufficio Tributi,

ha ragione: tre solleciti senza risposta non sono accettabili e ce ne
scusiamo. Abbiamo verificato la sua partecipazione al corso Bilancio e
contabilità degli enti locali e l'attestato le sarà inviato a questo
indirizzo entro [data, prima di venerdì].

Se non dovesse riceverlo, può scrivere direttamente a [contatto responsabile].

Cordiali saluti,
Segreteria Promo PA""", size=13,
    notes="Due segnaposto: la data e il contatto. Il modello non li conosce e non li inventa. L'operatore li riempie in dieci secondi.")

bullets_foto("Demo: perché ci sono i segnaposto", [
    "Una data inventata è un danno; un buco visibile è dieci secondi di lavoro",
    "La regola sta nel prompt di sistema, non nella buona volontà del modello",
    "È il modo più semplice per rendere visibile il confine di quello che il sistema sa",
], FOTO + "b50_puzzle.jpg", size=18,
    notes="Aneddoto da raccontare: la risposta piu' pericolosa non e' quella sbagliata, e' quella verosimile. I segnaposto rendono l'ignoranza visibile.")

foto_piena("Dodici richieste vere", FOTO + "hd_dodici.jpg",
           caption="Dodici richieste arrivate davvero alla Fondazione, anonimizzate, importate nel repository.",
           notes="Qui si cambia registro: fino a ora i dati erano sintetici. Da qui in poi sono i loro. Vale la pena dirlo esplicitamente, cambia l'attenzione della sala.")

bullets_foto("Cosa ci hanno insegnato le dodici 1/3", [
    "Dieci su dodici hanno corso 'nessuno': il catalogo vero non è quello della tassonomia di prova",
    "Lavoro agile, cyber security, OIV, performance, società partecipate: non c'erano",
    "Non è un errore del modello. È la tassonomia che va riscritta sul catalogo vostro",
], FOTO + "b52_scomparti.jpg", size=17,
    notes="Questo e' il risultato piu' utile della giornata e non era previsto. La tassonomia sintetica era plausibile e sbagliata.")

bullets_foto("Cosa ci hanno insegnato le dodici 2/3", [
    "Tre richieste su dodici sono commerciali: preventivi, sconti, codici MEPA",
    "La tassonomia non ha una voce per il commerciale, e le abbiamo messe sotto fatturazione",
    "È l'approssimazione migliore disponibile, ed è comunque sbagliata",
], FOTO + "b53_ricevute.jpg", size=17,
    notes="Domanda per la sala: chi risponde oggi a una richiesta di sconto? Se la risposta e' 'dipende', avete trovato la prossima voce della tassonomia.")

bullets_foto("Cosa ci hanno insegnato le dodici 3/3", [
    "Una richiesta su dodici non è una richiesta: è l'avviso automatico di una casella dismessa",
    "Il modello l'ha classificata come 'informazioni' e l'ha mandata alla segreteria",
    "Un sistema che smista tutto smista anche il rumore. Serve una voce per buttare via",
], FOTO + "b54_filtro.jpg", size=17,
    notes="RE12. E' l'argomento migliore a favore della voce 'altro' e della coda umana: qualcuno deve poter dire 'questa non e' una richiesta'.")

# ---------- 4. Limiti ----------
section("4. Limiti e rischi", "15 minuti", foto=FOTO + "sez4_limiti.jpg",
        notes="Senza giri di parole. Questa parte serve a evitare che qualcuno compri qualcosa a ottobre sull'onda dell'entusiasmo.")

bullets_foto("L'accuratezza si misura, non si stima", [
    "Serve un campione etichettato a mano: 100-200 richieste, due persone, poi confronto",
    "Dove le due persone non concordano, il problema è la tassonomia, non il modello",
    "Quello che segue è misurato sul repository, oggi, con quaranta richieste",
], FOTO + "b56_calibro.jpg", size=17,
    notes="Le quaranta sono: 19 sintetiche, 12 reali della Fondazione, 9 estratte da chat Zoom e questionari. Poche, ma vere e verificabili.")

image_full("I numeri di questa esecuzione", FIG + "accuratezza.png",
           caption="gpt-5-mini, 40 richieste, 10 settembre 2026. Etichette di riferimento scritte a mano.",
           notes="Il corso al 100 per cento e' meno bello di quel che sembra: dieci richieste su dodici erano 'nessuno', ed e' facile indovinare. L'urgenza al 72,5 e' il dato onesto.")

image_full("Il modello è sicuro anche quando sbaglia", FIG + "confidenza.png",
           caption="Nessuna richiesta è scesa sotto la soglia. Sedici su quaranta avevano almeno un campo sbagliato.",
           notes="Questa e' la slide piu' importante della giornata. Fermarsi. La soglia di confidenza a 0,7 non e' scattata nemmeno una volta: la confidenza minima dichiarata e' stata 0,75. Il punto di controllo automatico, da solo, non ha protetto niente.")

bullets_foto("Cosa dice quel grafico", [
    "La confidenza dichiarata non è una probabilità: è un numero che il modello sceglie",
    "Le crocette rosse stanno a destra della soglia quanto i pallini blu",
    "La coda umana va tarata sui dati, non sul numero che il modello si autoassegna",
], FOTO + "b59_bussola.jpg", size=18,
    notes="Come si tara davvero: si prende il campione etichettato, si ordina per confidenza e si guarda dove gli errori si diradano. Se non si diradano mai, la confidenza non serve e il controllo deve essere un altro.")

image_full("Dove si concentrano gli errori", FIG + "errori_campi.png",
           caption="Undici errori su quaranta sul campo urgenza, nessuno sul corso.",
           notes="L'urgenza e' il campo piu' soggettivo e il piu' sbagliato. Non e' un caso: e' quello dove anche due persone della segreteria non sarebbero d'accordo.")

bullets_foto("Urgenza: perché è il campo peggiore", [
    "La regola dice: alta solo con vincolo di tempo esplicito o blocco di accesso",
    "Ma 'avrei urgentemente bisogno' è un vincolo esplicito o solo un tono?",
    "Finché non lo decidete voi, il modello sceglie, e sceglie ogni volta in modo diverso",
], FOTO + "b61_sveglia.jpg", size=18,
    notes="RE10 dice 'avrei urgentemente bisogno': noi l'abbiamo etichettata alta, il modello media. Nessuno dei due ha torto. Manca la regola.")

bullets_foto("Quando sbaglia il modello e quando la tassonomia", [
    "Se due persone della segreteria darebbero risposte diverse, non è colpa del modello",
    "Se tutti darebbero la stessa risposta e il modello no, è colpa del modello",
    "Il test costa venti minuti e si fa su venti richieste, prima di comprare qualsiasi cosa",
], FOTO + "b62_bivio.jpg", size=17,
    notes="E' la diagnosi differenziale piu' utile che si portano a casa oggi. Scriverla alla lavagna.")

bullets_foto("La coda umana, in pratica", [
    "Qualcuno la apre ogni mattina, o non è un controllo: è un secondo arretrato",
    "Ogni richiesta rivista è un'etichetta gratis: rientra nel campione e migliora la misura",
    "Se la coda resta vuota per una settimana, la soglia è tarata male",
], FOTO + "hd_coda.jpg",
    notes="Il secondo punto e' quello che nessuno sfrutta: la coda umana e' una macchina che produce dati etichettati mentre lavora. Vale la pena progettarla per quello.")

bullets_foto("Cosa può andare storto", [
    "Il modello cambia versione e i numeri si spostano, senza che nessuno se ne accorga",
    "Un caso nuovo entra in una categoria vecchia perché non ce n'è una giusta",
    "La coda umana si riempie, nessuno la guarda, e diventa un secondo arretrato",
], FOTO + "li_storto.jpg",
    notes="Il terzo e' il fallimento piu' comune e il meno raccontato: la coda da verificare che diventa il posto dove le richieste vanno a morire.")

bullets_foto("Da qui in poi non parlo io", [
    "Sono un informatico, non un avvocato. Quello che segue è contesto, non un parere",
    "Nessuna slide di oggi dice se una norma si applica a voi: quella è una qualificazione giuridica",
    "Quello che posso darvi sono le domande giuste, e dove andare a cercare",
], FOTO + "pi_scritto.jpg", size=16,
    notes="Trenta secondi, detti guardando la sala. Serve a voi e serve a me. Da qui in avanti il registro cambia: non 'ecco cosa vale', ma 'ecco cosa chiedere a chi di dovere'. Se in aula arriva una domanda del tipo 'quindi possiamo farlo?', la risposta e' sempre la stessa: e' esattamente la domanda da mettere per iscritto a chi vi segue sul legale.")

bullets_foto("Dati personali: quello che riguarda il progetto", [
    "Le richieste contengono nomi, enti e a volte situazioni personali (malattia, contenziosi)",
    "Il testo di una richiesta esce dalla Fondazione e arriva a un fornitore esterno: è una scelta, e va decisa",
    "Al modello serve il testo della richiesta, non lo storico del mittente: la pipeline passa solo il testo",
], FOTO + "gag_timbri.jpg", size=16,
    notes="Questa slide parla di progettazione, non di conformita'. Il terzo punto e' una scelta tecnica che avete gia' fatta bene nel documento dei dodici esempi: nomi e codici sostituiti da X. Vale la pena dirlo.")

bullets_foto("Chi risponde delle decisioni sui dati", [
    "Prima della tecnologia serve una persona che decida se quel testo può uscire dalla Fondazione",
    "Se oggi quella casella è vuota, riempirla è il passo zero del pilota",
    "Chi debba riempirla, e con quale titolo, non lo decido io in questa slide",
], FOTO + "an_fornitore.jpg", size=16,
    notes="Detto senza drammi e senza diagnosi: non sto dicendo che siete inadempienti, sto dicendo che serve un nome accanto a una decisione. Se poi quel nome debba essere un responsabile della protezione dei dati e' una domanda per un legale.")

bullets_foto("Il perimetro, senza conclusioni", [
    "Ci sono tre testi di cui sentirete parlare: l'AI Act europeo, la legge italiana sull'IA, le linee guida AgID",
    "Nel 2026 le date si sono mosse più di una volta: qualunque cosa vi dica oggi va riverificata sul testo ufficiale",
    "Nessuno di questi testi nomina l'helpdesk di una fondazione formativa: la qualificazione la fa un giurista",
], FOTO + "b68_recinto.jpg", size=16,
    notes="Non leggere numeri di articolo dalla slide. Se qualcuno li chiede, stanno in docs/domande_legali.md con i link alle fonti e la data di consultazione, e con scritto che il consolidato non e' stato letto direttamente. Il messaggio della slide e' uno solo: il perimetro esiste, non lo traccio io.")

bullets_foto("Le sei domande da mettere per iscritto 1/2", [
    "Il sistema che vogliamo costruire rientra fra quelli soggetti a obblighi rafforzati? Su quale base?",
    "Dobbiamo dire a chi ci scrive che una parte del processo usa AI? In quale momento e con quali parole?",
    "Con quale base giuridica trattiamo il testo delle richieste, e dove va scritta?",
], FOTO + "b69_macchina.jpg", size=16,
    notes="Sono domande, non risposte. Suggerire di mandarle via e-mail a chi le deve firmare: una domanda scritta produce una risposta scritta, e una risposta scritta protegge chi la riceve.")

bullets_foto("Le sei domande da mettere per iscritto 2/2", [
    "Il fornitore del modello che ruolo assume rispetto ai nostri dati, e cosa dobbiamo firmare con lui?",
    "Serve una valutazione d'impatto prima di partire? Se sì, chi la redige e chi la firma?",
    "Ogni umano che togliete dal ciclo allunga questa lista: cosa cambierebbe se la risposta partisse da sola?",
], FOTO + "b70_cassetta.jpg", size=16,
    notes="L'ultimo punto e' il ponte con il resto della lezione: l'invio umano non e' solo una scelta di qualita', e' anche quello che tiene corta la lista delle domande. Non dire perche', dire solo che la lista si allunga.")

bullets_foto("Quanto costa davvero", [
    "Le chiamate al modello, su questi volumi, sono la voce più piccola del conto",
    "Il costo vero è la persona che ogni trimestre guarda gli errori e aggiorna il file",
    "Chiedete al fornitore il costo per mille richieste, e il costo di una modifica",
], FOTO + "li_costi.jpg",
    notes="Non mettere cifre: dipendono dal modello e cambiano ogni pochi mesi. [NEEDS SOURCE] per un costo per mille richieste da citare in aula, se lo si vuole.")

bullets_foto("Manutenzione", [
    "La tassonomia invecchia: nuovi corsi, nuovi ruoli. Qualcuno la possiede e l'aggiorna",
    "Il modello cambia: ogni cambio di versione si rimisura sul campione etichettato",
    "La coda umana è il termometro: se supera il 30%, si rivede la tassonomia, non il modello",
], FOTO + "gag_pianta.jpg",
    notes="La pianta secca accanto alla scrivania ordinata: e' esattamente cosa succede a una tassonomia che nessuno possiede.")

# ---------- 5. Pilota ----------
section("5. Dal caso d'uso al pilota", "15 minuti", foto=FOTO + "sez5_pilota.jpg",
        notes="docs/pilota_4_passi.md ha le stesse informazioni con qualche dettaglio in più. Da qui in poi si parla di cose da fare lunedì.")

image_top_bullets("Passo 1: audit del flusso attuale", FIG + "quattro_passi.png", [
    "Contare le richieste per canale per quattro settimane, telefonate comprese",
    "Cronometrare il tempo dalla ricezione alla prima risposta, per tipologia",
    "Costo: un foglio e la costanza di compilarlo. Nessuna tecnologia",
], img_h=1950000,
    notes="Una o due settimane. Il lavoro e' della segreteria, non di un consulente. Senza questi numeri, dopo non saprete dire se e' migliorato.")

image_top_bullets("Passo 2: tassonomia e dati etichettati", FIG + "quattro_passi.png", [
    "Scrivere il file con i corsi veri e i ruoli veri, non quelli di esempio",
    "Etichettare a mano 100-200 richieste storiche, due persone in parallelo",
    "Dove le due persone non concordano, si riscrive la voce. È il lavoro vero",
], img_h=1950000,
    notes="Due settimane. E' il passo che tutti vogliono saltare ed e' quello che decide se il resto funziona. Oggi ne abbiamo etichettate 40 e gia' si vede dove la tassonomia scricchiola.")

image_top_bullets("Passo 3: prototipo a basso codice", FIG + "quattro_passi.png", [
    "Contenitore unico, connettori dalle caselle, classificatore, coda da verificare",
    "Un GPT personalizzato basta per cominciare: il codice serve quando volete misurare",
    "L'invio resta umano per tutto il pilota, senza eccezioni",
], img_h=1950000,
    notes="Due-quattro settimane. Chi lo fa: qualcuno in Fondazione che sappia gia' usare lo strumento scelto. Se non c'e', il pilota va rimandato o va comprata la competenza.")

image_top_bullets("Passo 4: misurare un mese, poi decidere", FIG + "quattro_passi.png", [
    "Le stesse metriche del passo 1, sugli stessi canali, dopo un mese di uso",
    "Tre esiti possibili: estendere, correggere, fermare. Tutti e tre legittimi",
    "La decisione la prende la direzione, con i numeri davanti",
], img_h=1950000,
    notes="Un mese. La cosa piu' importante e' aver deciso le soglie prima, altrimenti si guarda il risultato e si decide che va bene comunque.")

image_full("Chi fa cosa", FIG + "ruoli.png",
           caption="Nessuna riga dice 'consulente esterno'. Se ne serve uno, sta accanto, non al posto.",
           notes="Se la Fondazione non ha nessuno per il passo 3, e' un'informazione preziosa: si compra quel pezzo, non tutto il progetto.")

image_full("Il calendario", FIG + "calendario.png",
           caption="Tre mesi dall'inizio alla decisione. I primi due passi si sovrappongono poco.",
           notes="Tre mesi e' realistico se nessuno lavora a tempo pieno. Comprimere il passo 2 e' l'errore classico.")

bullets_foto("Metriche decise prima di partire", [
    "Accuratezza per campo sul campione etichettato, non a sensazione",
    "Tempo alla prima risposta e ore di segreteria a settimana: prima e dopo",
    "Quota di richieste in coda umana, e quante di quelle vengono davvero riviste",
], FOTO + "pi_metro.jpg",
    notes="Il terzo punto nasce dall'errore visto prima: una coda che nessuno guarda non e' un controllo.")

bullets_foto("Le tre soglie, spiegate", [
    "Sotto l'80% su tipologia e operatore il sistema fa perdere tempo invece di darne",
    "Sopra il 30% di coda umana la tassonomia è ambigua, non il modello è scarso",
    "Se il tempo alla prima risposta non scende, il collo di bottiglia era altrove",
], FOTO + "b81_pesi.jpg", size=17,
    notes="Sono punti di partenza, non standard: vanno fissati dalla Fondazione prima di vedere i risultati. Oggi, su 40 richieste, tipologia e operatore stanno all'85 per cento e i quattro campi insieme al 60.")

bullets_foto("Strumenti a basso codice", [
    "ChatGPT Business: GPT personalizzato o Progetto con la tassonomia come file. Zero codice",
    "Connettori: Make, Zapier, n8n, oppure le automazioni di Microsoft 365 o Google Workspace già in uso",
    "Il codice del repository serve quando volete misurare e integrare per davvero",
], FOTO + "b82_cassetta_attrezzi.jpg", size=17,
    notes="Non fare pubblicità a un fornitore. Il criterio: prendete lo strumento che qualcuno in Fondazione sa già usare.")

foto_piena("Il report che nessuno legge", FOTO + "gag_stampante.jpg",
           caption="Una metrica che non cambia una decisione è carta.",
           notes="Da usare per chiudere il discorso metriche: per ogni numero che decidete di misurare, dite in anticipo quale decisione cambierebbe. Se non ce n'e' una, non misuratelo.")

bullets_foto("Cosa mettere per iscritto prima di partire", [
    "Chi possiede la tassonomia e ogni quanto la rivede",
    "Chi guarda la coda umana, quando, e cosa fa se cresce",
    "Le risposte alle sei domande del blocco 4, date per iscritto da chi ha titolo per darle",
], FOTO + "gag_pulsante.jpg",
    notes="Sono tre righe, non un documento. Ma senza queste tre righe il pilota dipende dalla buona volonta' di una persona, e finisce quando quella persona cambia ufficio.")

bullets_foto("Quello che il piano non prevede", [
    "Una persona chiave in ferie proprio nelle due settimane di etichettatura",
    "Il fornitore che cambia versione del modello mentre state misurando",
    "Mettete due settimane di margine e decidete in anticipo chi sostituisce chi",
], FOTO + "gag_ombrello.jpg",
    notes="Nessun piano di tre mesi regge senza margine. Meglio dichiararlo adesso che spiegarlo a dicembre.")

bullets_foto("Fermarsi è un esito legittimo", [
    "Il pilota serve a scoprire a basso costo se conviene, non a dimostrare che conviene",
    "Se dopo un mese i numeri non si muovono, si ferma e si è imparato qualcosa",
    "Il costo di fermarsi dopo tre mesi è mille volte più basso di quello di una piattaforma",
], FOTO + "pi_fermarsi.jpg",
    notes="Da dire con convinzione: dare il permesso di fermarsi e' la cosa che rende un pilota davvero un pilota.")

# ---------- 6. Esercitazione ----------
section("6. Esercitazione in gruppi", "20 minuti più 10 di restituzione", foto=FOTO + "sez6_esercitazione.jpg",
        notes="Gruppi da 3-4. Griglia stampata in esercitazione/griglia_scomposizione.md. Passare fra i tavoli, non restare davanti.")

bullets_foto("Consegna", [
    "Scegliete un processo della Fondazione: rendicontazione, rassegna stampa, iscrizioni, o uno vostro",
    "Scomponetelo con la griglia: un passo per riga, input, decisione, output",
    "Per ogni passo: integrazione, AI o umano?",
], FOTO + "c36_consegna.jpg",
    notes="Insistere: un passo per riga anche se sembra banale. La banalità è dove si nascondono le integrazioni.")

image_full("La griglia da compilare", FIG + "griglia_vuota.png",
           caption="Sei righe bastano. Se ne servono di più, il processo era due processi.",
           notes="Distribuire la griglia stampata. Dare venti minuti veri e avvisare a cinque minuti dalla fine.")

image_full("La griglia, compilata sull'helpdesk", FIG + "griglia.png",
           caption="Tre passi su sette con AI. Due restano umani. Il resto è integrazione o regola.",
           notes="Esempio da usare come modello. Chiedere ai gruppi di arrivare a una tabella così, non più bella.")

bullets_foto("Tre processi candidati", [
    "Rendicontazione di un progetto: documenti eterogenei, la firma resta umana",
    "Rassegna stampa e bandi: filtrare con criteri scritti da voi, decidere resta a voi",
    "Iscrizioni ai corsi: estrarre dati da testi liberi e segnalare le incongruenze",
], FOTO + "es_processi.jpg",
    notes="Dettagli in esercitazione/processi_candidati.md. Un processo loro e' sempre meglio di uno dei tre.")

foto_piena("Mentre lavorate", FOTO + "gag_cane.jpg",
           caption="Venti minuti. Nessuno vi guarda le spalle.",
           notes="Slide da lasciare proiettata durante l'esercitazione. Serve a togliere pressione e a segnare il tempo.")

bullets_foto("Restituzione: tre domande per gruppo", [
    "Qual è il passo che, automatizzato, fa risparmiare più tempo?",
    "Qual è il passo dove un errore del sistema costa di più?",
    "Quale numero misurereste dopo un mese per decidere se continuare?",
], FOTO + "c38_restituzione.jpg",
    notes="Due minuti a gruppo. Scrivere le risposte alla lavagna: di solito il passo che vale di più e quello che rischia di più sono vicini.")

bullets_foto("Tre errori da evitare", [
    "Automatizzare prima di misurare",
    "Togliere l'umano troppo presto",
    "Comprare la piattaforma prima di aver definito il processo",
], FOTO + "c39_errori.jpg", size=19,
    notes="Chiusura. Sono le tre righe della tesi iniziale, al contrario.")

bullets_foto("Che cosa vi portate a casa", [
    "Un repository con dati, codice, griglia e queste slide, che gira anche senza chiave API",
    "Tre domande da fare a un fornitore, e una diagnosi per capire chi ha sbagliato",
    "Una tassonomia da riscrivere sui vostri corsi veri: è il compito per lunedì",
], FOTO + "b95_borsa.jpg", size=17,
    notes="Il terzo punto e' il vero compito. Dieci minuti a settimana per un mese e il passo 2 e' fatto.")

bullets_foto("Materiali e contatti", [
    "Codice, dati, griglia e slide: github.com/<da completare>/promopa-agenti-ai",
    "Tutto gira anche senza chiave API (modalità 'mock') per provare la forma dell'output",
    "fabrizio.silvestri@uniroma1.it",
], FOTO + "b96_taccuino.jpg", size=18,
    notes="ATTENZIONE: il repository non è ancora pubblicato. Sostituire <da completare> con l'account GitHub prima della lezione, poi rigenerare il deck. Il comando di push è in HANDOFF.md.")

section("Grazie", "Domande?", foto=FOTO + "sez7_grazie.jpg", notes="")

prs.save(OUT)
print("salvato", OUT, "slide:", len(prs.slides))
