"""Importa gli esempi reali (anonimizzati) dal docx della Fondazione.

Il documento è una lista di richieste. Lo script prende ogni paragrafo
non vuoto (o ogni riga di tabella) come una richiesta, prova a riconoscere
il canale da parole chiave e scrive data/reali/richieste_reali.jsonl.
Il canale e il mittente vanno poi controllati a mano: è il primo esempio
di "l'umano etichetta, il modello applica".

Uso:
    python -m helpdesk_agent.importa_docx "data/reali/Richieste a Promo PA Fondazione.docx"
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

from .schema import RADICE

CANALI = [
    ("zoom_chat", r"zoom|chat"),
    ("moodle_questionario", r"questionario|gradimento"),
    ("moodle_messaggio", r"moodle"),
    ("form_sito", r"modulo|sito|form"),
]


def indovina_canale(testo: str) -> str:
    t = testo.lower()
    for canale, pattern in CANALI:
        if re.search(pattern, t):
            return canale
    return "email"


def estrai_paragrafi(percorso: Path) -> list[str]:
    doc = Document(str(percorso))
    blocchi = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tabella in doc.tables:
        for riga in tabella.rows:
            celle = [c.text.strip() for c in riga.cells if c.text.strip()]
            if celle:
                blocchi.append(" | ".join(celle))
    return blocchi


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    sorgente = Path(sys.argv[1])
    uscita = RADICE / "data" / "reali" / "richieste_reali.jsonl"
    uscita.parent.mkdir(parents=True, exist_ok=True)
    blocchi = estrai_paragrafi(sorgente)
    n = 0
    with uscita.open("w", encoding="utf-8") as f:
        for i, testo in enumerate(blocchi, 1):
            if len(testo) < 25:  # titoli, numerazioni, righe vuote di fatto
                continue
            n += 1
            f.write(json.dumps(dict(
                id=f"REALE-{n:02d}", canale=indovina_canale(testo),
                sorgente="da verificare", mittente="Anonimo", testo=testo,
            ), ensure_ascii=False) + "\n")
    print(f"Scritte {n} richieste in {uscita}. Controlla canale e mittente a mano.")


if __name__ == "__main__":
    main()
